#!/usr/bin/env python3
"""Build the BREADCRUMBS pitch deck as .pptx — importable into Figma Slides.

Why PowerPoint and not the Figma API: Figma Slides imports .pptx directly, and
.pptx carries speaker notes natively. That routes around the Starter-plan MCP
call limit entirely, and the same file opens in Keynote, PowerPoint and Google
Slides.

Content is the Team Bible (docs/team-bible.md, AUTHORITATIVE). Speaker notes are
the same 15 blocks as scripts/figma-speaker-notes.js, verbatim, so the two decks
stay in sync.

    python3 deck/build_deck.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "deck" / "breadcrumbs-deck.pptx"
LOGO = ROOT / "brand" / "breadcrumbs-logo-dark.png"

# --- palette: the product's, from globe/food3d.html and brand/ ---------------
ABYSS = RGBColor(0x06, 0x0A, 0x12)
PANEL = RGBColor(0x0B, 0x11, 0x1C)
LINE = RGBColor(0x1B, 0x27, 0x40)
CYAN = RGBColor(0x3F, 0xB6, 0xC9)
INK = RGBColor(0xE8, 0xED, 0xF4)
MUT = RGBColor(0x6B, 0x7A, 0x95)
CRIT = RGBColor(0xE5, 0x48, 0x4D)
AMBER = RGBColor(0xE0, 0xA9, 0x5A)

# Deliberately not Inter. Both are Google Fonts; if absent the renderer
# substitutes, which is why the fallbacks are named in deck/README.md.
HEAD = "Archivo"
MONO = "IBM Plex Mono"

W, H = Inches(13.333), Inches(7.5)
ML = Inches(0.85)                       # left margin
CW = W - ML * 2                         # content width

# The Bible's §C site table uses an anonymised illustrative portfolio
# (Site 03/05/07/09/11, masked CAMIS, crit_violations 14/9/6/4/2). The live
# backend returns real named sites with 11/10/10/8/8. They are different
# datasets, not a bug. The embedded speaker notes quote "Site 05 (LES, 14)" and
# "Site 11 (Harlem, 9)" verbatim, so the slide follows the Bible to stay
# internally consistent with what gets read aloud. Flip to "backend" if the
# deck should instead mirror what the demo screen shows.
SITE_TABLE = "bible"

SITES_BIBLE = [("Site 05 — LES", "14", "+10", "Elevated"),
               ("Site 11 — Harlem", "9", "+5", "Above median"),
               ("Site 03 — Midtown", "6", "+2", "Slightly above"),
               ("Site 07 — Williamsburg", "4", "0", "At median"),
               ("Site 09 — LIC", "2", "−2", "Below median")]
SITES_BACKEND = [("JUST SALAD — Manhattan", "11", "+7", "Elevated"),
                 ("JUST SALAD — Brooklyn", "10", "+6", "Elevated"),
                 ("SALAD DON — Manhattan", "10", "+6", "Elevated"),
                 ("SWEETGREEN — Manhattan", "8", "+4", "Above median"),
                 ("JUST SALAD — Manhattan", "8", "+4", "Above median")]
SITES = SITES_BIBLE if SITE_TABLE == "bible" else SITES_BACKEND


# --------------------------------------------------------------------------
# primitives
# --------------------------------------------------------------------------
def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])          # blank
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = ABYSS
    return s


def rect(slide, x, y, w, h, fill=None, edge=None, edge_w=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if edge is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = edge
        sh.line.width = Pt(edge_w)
    sh.text_frame.word_wrap = True
    return sh


def rule(slide, x, y, w, color=LINE, weight=0.75):
    return rect(slide, x, y, w, Pt(weight), fill=color)


def text(slide, x, y, w, blocks, h=Inches(1.0), anchor=MSO_ANCHOR.TOP):
    """blocks: list of dicts -> t, size, font, color, bold, align, space, lead, caps"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    first = True
    for b in blocks:
        body = b.get("t", "")
        if b.get("caps"):
            body = body.upper()
        # One paragraph per line. A "\n" inside a run renders as a soft break
        # that renderers align and space inconsistently -- centred lines drift
        # left and leading collapses into full paragraph gaps.
        for j, ln in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = b.get("align", PP_ALIGN.LEFT)
            if j == 0 and b.get("space"):
                p.space_before = Pt(b["space"])
            if b.get("lead"):
                p.line_spacing = b["lead"]
            r = p.add_run()
            r.text = ln
            f = r.font
            f.name = b.get("font", HEAD)
            f.size = Pt(b.get("size", 16))
            f.bold = b.get("bold", False)
            f.color.rgb = b.get("color", INK)
            if b.get("track"):
                # python-pptx has no letter-spacing API; set it on the raw XML.
                f._rPr.set("spc", str(int(b["track"] * 100)))
    return box


def chrome(slide, section, number, title=None, sub=None):
    """Eyebrow + slide number + optional title block. Returns next free y."""
    text(slide, ML, Inches(0.46), CW,
         [{"t": section, "size": 10, "font": MONO, "color": MUT,
           "caps": True, "track": 2.2}], h=Inches(0.3))
    text(slide, W - ML - Inches(1.2), Inches(0.46), Inches(1.2),
         [{"t": f"{number:02d}", "size": 10, "font": MONO, "color": LINE,
           "align": PP_ALIGN.RIGHT, "track": 1.2}], h=Inches(0.3))
    y = Inches(1.02)
    if title:
        text(slide, ML, y, CW,
             [{"t": title, "size": 40, "bold": True, "lead": 1.06}], h=Inches(1.0))
        y = Inches(1.92)
        if sub:
            text(slide, ML, y, CW,
                 [{"t": sub, "size": 15, "font": MONO, "color": CYAN, "lead": 1.4}],
                 h=Inches(0.5))
            y = Inches(2.52)
        else:
            y = Inches(2.20)
    return y


def kv_rows(slide, x, y, w, rows, cols, head=None, row_h=Inches(0.42)):
    """Hand-drawn table: full control, no PowerPoint table styling to fight."""
    cy = y
    if head:
        cx = x
        for label, frac in zip(head, cols):
            text(slide, cx, cy, Emu(int(w * frac)),
                 [{"t": label, "size": 9, "font": MONO, "color": MUT,
                   "caps": True, "track": 1.6}], h=Inches(0.25))
            cx += Emu(int(w * frac))
        cy += Inches(0.32)
        rule(slide, x, cy, w, LINE)
        cy += Inches(0.16)

    for row in rows:
        cx = x
        for cell, frac in zip(row, cols):
            body, colour, mono, bold = cell
            text(slide, cx, cy, Emu(int(w * frac)) - Inches(0.15),
                 [{"t": body, "size": 13 if not mono else 13,
                   "font": MONO if mono else HEAD,
                   "color": colour, "bold": bold, "lead": 1.25}],
                 h=row_h)
            cx += Emu(int(w * frac))
        cy += row_h
        rule(slide, x, cy - Inches(0.09), w, PANEL)
    return cy


def cell(t, colour=INK, mono=False, bold=False):
    return (t, colour, mono, bold)


def node(slide, x, y, w, h, label, detail, accent=LINE, label_col=INK):
    rect(slide, x, y, w, h, fill=PANEL, edge=accent)
    text(slide, x + Inches(0.16), y + Inches(0.14), w - Inches(0.32),
         [{"t": label, "size": 13, "bold": True, "color": label_col},
          {"t": detail, "size": 9.5, "font": MONO, "color": MUT,
           "space": 4, "lead": 1.3}], h=h - Inches(0.28))


def arrow(slide, x, y, w):
    text(slide, x, y, w, [{"t": "→", "size": 17, "color": CYAN,
                           "align": PP_ALIGN.CENTER}], h=Inches(0.3))


# --------------------------------------------------------------------------
# slides
# --------------------------------------------------------------------------
def s01_title(prs):
    s = new_slide(prs)
    if LOGO.exists():
        lw = Inches(4.3)
        s.shapes.add_picture(str(LOGO), int((W - lw) / 2), Inches(1.02), width=lw)
    text(s, ML, Inches(3.30), CW,
         [{"t": "Our food supply is a black box.", "size": 42, "bold": True,
           "align": PP_ALIGN.CENTER, "lead": 1.1}], h=Inches(0.8))
    text(s, ML, Inches(4.16), CW,
         [{"t": "Follow the crumbs back to the source —\n"
                "and forward to every plate at risk.",
           "size": 19, "font": MONO, "color": CYAN,
           "align": PP_ALIGN.CENTER, "lead": 1.5}], h=Inches(1.1))
    rule(s, int((W - Inches(1.4)) / 2), Inches(5.60), Inches(1.4), LINE)
    text(s, ML, Inches(5.94), CW,
         [{"t": "always-on food traceability   ·   on-prem   ·   NVIDIA GB10",
           "size": 11, "font": MONO, "color": MUT,
           "align": PP_ALIGN.CENTER, "caps": True, "track": 1.8}], h=Inches(0.4))
    return s


def s02_today(prs):
    s = new_slide(prs)
    y = chrome(s, "Problem", 2, "How it works today",
               "distributor email  →  days later  →  spreadsheet hunt  →  phone tree")

    bw, gap = Inches(2.55), Inches(0.28)
    steps = [("Distributor email", "arrives days after the recall posts"),
             ("Spreadsheet hunt", "which locations bought the lot?"),
             ("Phone tree", "12, 50, or 200 kitchens, one at a time"),
             ("Action", "if it is still in time to matter")]
    x = ML
    for i, (label, detail) in enumerate(steps):
        node(s, x, y, bw, Inches(1.18), label, detail,
             accent=CRIT if i == 2 else LINE)
        x += bw
        if i < len(steps) - 1:
            arrow(s, x, y + Inches(0.38), gap)
            x += gap

    y2 = y + Inches(1.72)
    text(s, ML, y2, CW,
         [{"t": "The stakes are real money", "size": 11, "font": MONO,
           "color": MUT, "caps": True, "track": 1.8}], h=Inches(0.3))
    y2 += Inches(0.46)

    money = [("$0.04M–$1.1M", "median per-recall cost, per restaurant firm", AMBER),
             ("~$10M", "average direct recall cost for grocers (FMI/GMA)", AMBER),
             ("52%", "of major recalls exceed $10M in total impact", CRIT)]
    cwid = int((CW - Inches(0.6)) / 3)
    x = ML
    for big, cap, col in money:
        text(s, x, y2, Emu(cwid),
             [{"t": big, "size": 30, "bold": True, "color": col},
              {"t": cap, "size": 11, "font": MONO, "color": MUT,
               "space": 8, "lead": 1.4}], h=Inches(1.3))
        x += cwid + Inches(0.3)

    text(s, ML, Inches(6.62), CW,
         [{"t": "The reputational tail is far larger.", "size": 12,
           "font": MONO, "color": MUT}], h=Inches(0.34))
    return s


def s03_contrast(prs):
    s = new_slide(prs)
    chrome(s, "Problem", 3)
    half = int(CW / 2)

    text(s, ML, Inches(2.15), Emu(half),
         [{"t": "BREADCRUMBS", "size": 11, "font": MONO, "color": CYAN,
           "caps": True, "track": 2.0},
          {"t": "4 seconds", "size": 64, "bold": True, "color": INK,
           "space": 14, "lead": 1.0},
          {"t": "on a box in the back office,\nsupplier list never leaving the building",
           "size": 13, "font": MONO, "color": MUT, "space": 16, "lead": 1.5}],
         h=Inches(3.0))

    rect(s, ML + Emu(half) - Inches(0.02), Inches(2.15), Pt(0.75), Inches(2.3), fill=LINE)

    text(s, ML + Emu(half) + Inches(0.5), Inches(2.15), Emu(half) - Inches(0.5),
         [{"t": "Today", "size": 11, "font": MONO, "color": MUT,
           "caps": True, "track": 2.0},
          {"t": "2 days", "size": 64, "bold": True, "color": MUT,
           "space": 14, "lead": 1.0},
          {"t": "email, spreadsheets, and a phone tree\nacross every location",
           "size": 13, "font": MONO, "color": MUT, "space": 16, "lead": 1.5}],
         h=Inches(3.0))

    text(s, ML, Inches(5.62), CW,
         [{"t": "That's the difference between a recall and an outbreak.",
           "size": 22, "bold": True, "color": CYAN, "align": PP_ALIGN.CENTER}],
         h=Inches(0.6))
    return s


def s04_local(prs):
    s = new_slide(prs)
    y = chrome(s, "Problem", 4, "Why it has to run local")

    text(s, ML, y, Emu(int(CW * 0.54)),
         [{"t": "A food business's supplier list, purchase volumes and margins "
                "are trade secrets.", "size": 20, "bold": True, "lead": 1.3},
          {"t": "No one POSTs that to a cloud LLM to get a recall alert.\n"
                "The alert isn't worth the leak.",
           "size": 14, "font": MONO, "color": MUT, "space": 18, "lead": 1.55},
          {"t": "Local isn't a constraint we're excusing —\n"
                "it's the only shape the customer will accept.",
           "size": 15, "font": MONO, "color": CYAN, "space": 22, "lead": 1.5}],
         h=Inches(3.4))

    bx = ML + Emu(int(CW * 0.60))
    bw = CW - Emu(int(CW * 0.60))
    rect(s, bx, y, bw, Inches(3.34), fill=PANEL, edge=AMBER)
    text(s, bx + Inches(0.28), y + Inches(0.30), bw - Inches(0.56),
         [{"t": "FSMA 204", "size": 26, "bold": True, "color": AMBER},
          {"t": "FDA Food Traceability Rule", "size": 11, "font": MONO,
           "color": MUT, "space": 6},
          {"t": "MANDATORY\nJULY 20, 2028", "size": 20, "bold": True,
           "color": INK, "space": 18, "lead": 1.2},
          {"t": "Lot-level traceability becomes a legal recordkeeping "
                "mandate — so this data is about to exist in every back "
                "office whether they like it or not.",
           "size": 10.5, "font": MONO, "color": MUT, "space": 16, "lead": 1.45}],
         h=Inches(2.1))
    return s


def s05_demo(prs):
    s = new_slide(prs)
    y = chrome(s, "Demo", 5, "Live demo",
               "a real NYC food business · 12 sites · their actual buildings")

    facts = [("12", "sites in the portfolio"),
             ("0", "bytes leaving the box"),
             ("29,310", "FDA recalls indexed locally"),
             ("17,204", "geocoded sites with violation counts")]
    cwid = int((CW - Inches(0.9)) / 4)
    x = ML
    for big, cap in facts:
        text(s, x, y, Emu(cwid),
             [{"t": big, "size": 34, "bold": True, "color": CYAN},
              {"t": cap, "size": 10.5, "font": MONO, "color": MUT,
               "space": 8, "lead": 1.4}], h=Inches(1.4))
        x += cwid + Inches(0.3)

    yb = y + Inches(1.86)
    rect(s, ML, yb, CW, Inches(1.42), fill=PANEL, edge=LINE)
    text(s, ML + Inches(0.30), yb + Inches(0.24), CW - Inches(0.6),
         [{"t": "Say it, don't wait to be asked", "size": 9.5, "font": MONO,
           "color": AMBER, "caps": True, "track": 1.6},
          {"t": "“The buildings and inspection data are real. This 12-site "
                "portfolio is an illustrative operator we modeled — I'll show "
                "you what's real underneath in a second.”",
           "size": 14, "color": INK, "space": 10, "lead": 1.45}],
         h=Inches(0.95))

    text(s, ML, Inches(6.40), CW,
         [{"t": "BREADCRUMBS is watching the FDA recall feed for them right "
                "now, and every byte is running on this box. No internet.",
           "size": 12, "font": MONO, "color": MUT, "lead": 1.4}], h=Inches(0.6))
    return s


def s06_recall(prs):
    s = new_slide(prs)
    y = chrome(s, "Demo", 6, "A Class I recall drops")

    rect(s, ML, y, CW, Inches(2.30), fill=PANEL, edge=CRIT)
    text(s, ML + Inches(0.34), y + Inches(0.26), CW - Inches(0.68),
         [{"t": "F-0757-2022", "size": 13, "font": MONO, "color": CRIT,
           "track": 1.4},
          {"t": "DOLE FRESH VEGETABLES INC — iceberg lettuce", "size": 27,
           "bold": True, "space": 10},
          {"t": "Listeria monocytogenes   ·   Class I   ·   firm state CA",
           "size": 14, "font": MONO, "color": MUT, "space": 12}],
         h=Inches(1.7))

    yb = y + Inches(2.62)
    half = int((CW - Inches(0.5)) / 2)
    text(s, ML, yb, Emu(half),
         [{"t": "What Class I means", "size": 10, "font": MONO, "color": MUT,
           "caps": True, "track": 1.8},
          {"t": "“A reasonable probability of\nserious adverse health "
                "consequences or death.”", "size": 17, "bold": True,
           "space": 12, "lead": 1.35},
          {"t": "21 CFR 7.3 — the FDA's most serious classification.",
           "size": 11.5, "font": MONO, "color": MUT, "space": 12}],
         h=Inches(2.0))

    text(s, ML + Emu(half) + Inches(0.5), yb, Emu(half),
         [{"t": "Provenance", "size": 10, "font": MONO, "color": MUT,
           "caps": True, "track": 1.8},
          {"t": "REAL", "size": 30, "bold": True, "color": CYAN, "space": 8},
          {"t": "openFDA food enforcement.\nReal recall, real number — "
                "one of 29,310 in the local corpus.",
           "size": 12.5, "font": MONO, "color": MUT, "space": 10, "lead": 1.45}],
         h=Inches(2.0))
    return s


def s07_trace(prs):
    s = new_slide(prs)
    y = chrome(s, "Demo", 7, "Trace it forward")

    text(s, ML, y, Emu(int(CW * 0.30)),
         [{"t": "5 of 12", "size": 62, "bold": True, "color": CRIT, "lead": 1.0},
          {"t": "sites received the lot.\nThe other seven are fine — and "
                "knowing which is the whole game.",
           "size": 13, "font": MONO, "color": MUT, "space": 14, "lead": 1.5}],
         h=Inches(2.6))

    tx = ML + Emu(int(CW * 0.34))
    tw = CW - Emu(int(CW * 0.34))
    rows = [(cell(n, INK), cell("Dole → iceberg", MUT, mono=True),
             cell("Receiving log", MUT, mono=True))
            for n, *_ in SITES]
    kv_rows(s, tx, y, tw, rows, [0.38, 0.34, 0.28],
            head=["Site", "Supplier link → this lot", "Verify against"],
            row_h=Inches(0.46))

    yb = Inches(5.30)
    rect(s, ML, yb, CW, Inches(1.42), fill=PANEL, edge=AMBER)
    text(s, ML + Inches(0.30), yb + Inches(0.22), CW - Inches(0.6),
         [{"t": "MODELED — say it in the same breath as the number",
           "size": 9.5, "font": MONO, "color": AMBER, "caps": True, "track": 1.5},
          {"t": "The supplier→site link is the one modeled layer — and that's "
                "the point. FSMA 204 is the rule that's supposed to make this "
                "traceable and doesn't yet. We model the gap the law is trying "
                "to close.", "size": 13, "space": 10, "lead": 1.45}],
         h=Inches(1.0))
    return s


def s08_signals(prs):
    s = new_slide(prs)
    y = chrome(s, "Demo", 8, "Two independent, real signals")

    half = int((CW - Inches(0.7)) / 2)

    text(s, ML, y, Emu(half),
         [{"t": "1 — Distributor recall history", "size": 11, "font": MONO,
           "color": CYAN, "caps": True, "track": 1.6}], h=Inches(0.3))
    firms = [(cell("Dole", INK, bold=True), cell("205", CRIT, mono=True, bold=True),
              cell("Extensive history", MUT, mono=True)),
             (cell("Fresh Express"), cell("110", AMBER, mono=True, bold=True),
              cell("High", MUT, mono=True)),
             (cell("Sysco"), cell("24", AMBER, mono=True, bold=True),
              cell("Moderate", MUT, mono=True)),
             (cell("Jetro", CYAN, bold=True), cell("0 of 1", CYAN, mono=True, bold=True),
              cell("Verified low record", MUT, mono=True))]
    kv_rows(s, ML, y + Inches(0.40), Emu(half), firms, [0.36, 0.22, 0.42],
            head=["Firm", "Class-I", "Read"], row_h=Inches(0.46))
    text(s, ML, y + Inches(3.02), Emu(half),
         [{"t": "REAL — openFDA firm-level recall counts", "size": 10,
           "font": MONO, "color": MUT}], h=Inches(0.3))

    x2 = ML + Emu(half) + Inches(0.7)
    text(s, x2, y, Emu(half),
         [{"t": "2 — Site handling risk", "size": 11, "font": MONO,
           "color": CYAN, "caps": True, "track": 1.6}], h=Inches(0.3))
    rows = [(cell(n), cell(v, CRIT if d.startswith("+1") else INK, mono=True, bold=True),
             cell(d, MUT, mono=True), cell(r, MUT, mono=True))
            for n, v, d, r in SITES[:4]]
    kv_rows(s, x2, y + Inches(0.40), Emu(half), rows, [0.40, 0.14, 0.14, 0.32],
            head=["Site", "Crit", "vs 4", "Read"], row_h=Inches(0.46))
    text(s, x2, y + Inches(3.02), Emu(half),
         [{"t": "REAL — DOHMH crit_violations (0–56, median 4)",
           "size": 10, "font": MONO, "color": MUT}], h=Inches(0.3))

    yb = Inches(5.56)
    rect(s, ML, yb, CW, Inches(1.30), fill=PANEL, edge=LINE)
    text(s, ML + Inches(0.30), yb + Inches(0.20), CW - Inches(0.6),
         [{"t": "0 of 1 is a VERIFIED LOW RECORD — never “clean”.",
           "size": 16, "bold": True, "color": CYAN},
          {"t": "An absence of matched records is not proof of a safer "
                "product; it can also mean lower recall visibility.   ·   "
                "Never say “C health grade” — grades composite vermin, "
                "facility, paperwork and temperature. We use crit_violations.",
           "size": 11, "font": MONO, "color": MUT, "space": 10, "lead": 1.45}],
         h=Inches(0.9))
    return s


def s09_output(prs):
    s = new_slide(prs)
    y = chrome(s, "Credibility", 9, "What the manager actually gets",
               "a document a health inspector, an insurer, or a plaintiff's "
               "attorney already recognises")

    lw = Emu(int(CW * 0.46))
    rect(s, ML, y, lw, Inches(3.44), fill=PANEL, edge=LINE)
    text(s, ML + Inches(0.28), y + Inches(0.24), lw - Inches(0.56),
         [{"t": "Recall Exposure & Action Report", "size": 15, "bold": True},
          {"t": "A   FDA recall facts — mapped 1:1 to openFDA\n"
                "B   Exposure — which sites received the lot\n"
                "C   Risk — two independent real signals\n"
                "D   Recommended action — deterministic\n"
                "E   Real vs. modeled provenance",
           "size": 12, "font": MONO, "color": INK, "space": 14, "lead": 1.85},
          {"t": "21 CFR 7.3 · classification\n21 CFR 7.49 · recall communication",
           "size": 10.5, "font": MONO, "color": CYAN, "space": 16, "lead": 1.5}],
         h=Inches(2.6))

    rx = ML + lw + Inches(0.5)
    rw = CW - lw - Inches(0.5)
    rect(s, rx, y, rw, Inches(2.24), fill=ABYSS, edge=LINE)
    text(s, rx + Inches(0.26), y + Inches(0.20), rw - Inches(0.52),
         [{"t": "Telegram / SMS outbox", "size": 9.5, "font": MONO,
           "color": MUT, "caps": True, "track": 1.6},
          {"t": "FDA CLASS I RECALL — action needed today.\n"
                "Dole iceberg lettuce (Listeria), recall F-0757-2022.\n"
                "5 of your 12 sites received matching lots: Midtown,\n"
                "LES, Williamsburg, LIC, Harlem.\n"
                "DO: hold & segregate now; verify lot codes; do not serve.\n"
                "Full report + swap: BC-2026-0822-0417 (PDF).",
           "size": 10.5, "font": MONO, "color": INK, "space": 10, "lead": 1.5}],
         h=Inches(1.8))

    text(s, rx, y + Inches(2.44), rw,
         [{"t": "Drafted, not dispatched.", "size": 15, "bold": True,
           "color": AMBER},
          {"t": "A human still hits go.", "size": 12, "font": MONO,
           "color": MUT, "space": 6}], h=Inches(0.8))

    text(s, ML, Inches(6.42), CW,
         [{"t": "A dated, on-prem-generated report showing what they knew, "
                "when, and what action was directed is contemporaneous "
                "evidence of a good-faith response.",
           "size": 12, "font": MONO, "color": MUT, "lead": 1.4}], h=Inches(0.6))
    return s


def s10_coverage(prs):
    s = new_slide(prs)
    y = chrome(s, "Credibility", 10, "Where we work — and where we don't",
               "volunteer this before a judge digs for it")

    rows = [
        (cell("Jif peanut butter, 2022", INK, bold=True), cell("YES", CYAN, mono=True, bold=True),
         cell("FDA Class-I lot-coded recall — exactly our trigger. 24 real Smucker records verified.", MUT, mono=True)),
        (cell("McDonald's / Taylor Farms\nonions, 2024", INK, bold=True), cell("YES", CYAN, mono=True, bold=True),
         cell("FDA-regulated produce, lot-coded — we map it to the sites that received the lot.", MUT, mono=True)),
        (cell("Yuma romaine, 2018", INK, bold=True), cell("PARTIAL", AMBER, mono=True, bold=True),
         cell("Commodity produce → a public advisory, not a lot-coded recall. No lot to trace.", MUT, mono=True)),
        (cell("Boar's Head deli meat, 2024", INK, bold=True), cell("NO", CRIT, mono=True, bold=True),
         cell("USDA FSIS jurisdiction — invisible to the openFDA food feed. Our labeled boundary.", MUT, mono=True)),
        (cell("Chipotle, 2015", INK, bold=True), cell("NO", CRIT, mono=True, bold=True),
         cell("No recall was ever issued and the ingredient was never identified — nothing to trigger on.", MUT, mono=True)),
    ]
    kv_rows(s, ML, y, CW, rows, [0.26, 0.11, 0.63],
            head=["Real US outbreak", "Relevant?", "Why"], row_h=Inches(0.64))

    yb = Inches(6.24)
    text(s, ML, yb, CW,
         [{"t": "~100% of FDA lot-coded enforcement recalls. A minority of the "
                "full universe of epidemiological outbreaks.",
           "size": 14, "bold": True, "color": CYAN, "lead": 1.35},
          {"t": "We sit downstream of CDC/FDA traceback and act the instant a "
                "recall posts.   ·   Real cadence: 24 NY-distributed Class-I "
                "recalls in 2025–26 — roughly 1–2 a month.",
           "size": 11, "font": MONO, "color": MUT, "space": 8, "lead": 1.4}],
         h=Inches(1.0))
    return s


def s11_instrument(prs):
    s = new_slide(prs)
    y = chrome(s, "Credibility", 11, "An instrument, not an oracle")

    text(s, ML, y, CW,
         [{"t": "“Unset the API key — what changes?”", "size": 20,
           "bold": True, "color": AMBER},
          {"t": "The narration sentence goes generic. Nothing else. Every red "
                "site, every risk score, the swap, the report — all "
                "deterministic Mongo aggregation. The LLM writes prose; it "
                "doesn't make the decision.",
           "size": 14, "font": MONO, "color": INK, "space": 12, "lead": 1.55}],
         h=Inches(1.8))

    yb = y + Inches(1.58)
    half = int((CW - Inches(0.6)) / 2)

    rect(s, ML, yb, Emu(half), Inches(2.44), fill=PANEL, edge=CYAN)
    text(s, ML + Inches(0.26), yb + Inches(0.20), Emu(half) - Inches(0.52),
         [{"t": "REAL", "size": 15, "bold": True, "color": CYAN, "track": 1.4},
          {"t": "29,310 FDA recalls\n"
                "17,204 geocoded sites with critical-violation counts\n"
                "Distributor recall histories\n"
                "LiDAR building geometry",
           "size": 12, "font": MONO, "color": INK, "space": 10, "lead": 1.5}],
         h=Inches(1.7))

    rect(s, ML + Emu(half) + Inches(0.6), yb, Emu(half), Inches(2.44),
         fill=PANEL, edge=AMBER)
    text(s, ML + Emu(half) + Inches(0.86), yb + Inches(0.20), Emu(half) - Inches(0.52),
         [{"t": "MODELED — and labeled", "size": 15, "bold": True,
           "color": AMBER, "track": 1.4},
          {"t": "Which supplier feeds which site\n"
                "The illustrative 12-site portfolio",
           "size": 12, "font": MONO, "color": INK, "space": 10, "lead": 1.5}],
         h=Inches(1.7))

    text(s, ML, Inches(6.46), CW,
         [{"t": "When a real food business plugs in their real supplier list, "
                "the modeled layer disappears.",
           "size": 15, "bold": True, "color": CYAN, "align": PP_ALIGN.CENTER}],
         h=Inches(0.6))
    return s


def s12_swarm(prs):
    s = new_slide(prs)
    y = chrome(s, "Credibility", 12, "Five stages. The LLM touches one.")

    bw = Inches(2.16)
    gap = Inches(0.26)

    node(s, ML, y + Inches(0.62), bw, Inches(1.10), "Watcher",
         "detects a new openFDA\nenforcement record")
    arrow(s, ML + bw, y + Inches(1.00), gap)

    x2 = ML + bw + gap
    node(s, x2, y, bw, Inches(1.16), "Tracer",
         "recall firm → suppliers,\nthen the sites they feed")
    node(s, x2, y + Inches(1.30), bw, Inches(1.16), "Risk",
         "crit_violations +\nrecall history")
    text(s, x2, y + Inches(2.54), bw,
         [{"t": "run concurrently", "size": 9.5, "font": MONO, "color": CYAN,
           "align": PP_ALIGN.CENTER}], h=Inches(0.3))
    arrow(s, x2 + bw, y + Inches(1.00), gap)

    x3 = x2 + bw + gap
    node(s, x3, y + Inches(0.62), bw, Inches(1.10), "Briefer",
         "site list, action,\nsupplier-swap recommendation")
    arrow(s, x3 + bw, y + Inches(1.00), gap)

    x4 = x3 + bw + gap
    node(s, x4, y + Inches(0.62), bw, Inches(1.10), "Comms",
         "wording only", accent=AMBER, label_col=AMBER)
    text(s, x4, y + Inches(1.82), bw,
         [{"t": "the only LLM stage", "size": 9.5, "font": MONO, "color": AMBER,
           "align": PP_ALIGN.CENTER}], h=Inches(0.3))

    yb = Inches(5.06)
    text(s, ML, yb, CW,
         [{"t": "Four of five stages are pure deterministic computation over "
                "real data — an orchestration of auditable steps, not a chain "
                "of LLM guesses.",
           "size": 15, "bold": True, "lead": 1.4},
          {"t": "MongoDB earns its place: 2dsphere $near for geospatial, "
                "$lookup/$group to roll up exposure at query time, and a "
                "change stream that watches. A spreadsheet is a snapshot; "
                "this is a watcher.",
           "size": 11.5, "font": MONO, "color": MUT, "space": 12, "lead": 1.45},
          {"t": "Model: Llama-3.1-Nemotron-Nano-4B via vLLM on localhost:8000",
           "size": 10.5, "font": MONO, "color": CYAN, "space": 10}],
         h=Inches(2.0))
    return s


def s13_business(prs):
    s = new_slide(prs)
    y = chrome(s, "Business", 13, "The hinge is a regulation",
               "and it writes our data for us")

    lw = Emu(int(CW * 0.40))
    rect(s, ML, y, lw, Inches(2.10), fill=PANEL, edge=AMBER)
    text(s, ML + Inches(0.28), y + Inches(0.26), lw - Inches(0.56),
         [{"t": "JULY 20, 2028", "size": 26, "bold": True, "color": AMBER},
          {"t": "FSMA 204 becomes mandatory. To comply, food businesses must "
                "generate lot codes and key data elements at every critical "
                "tracking event.",
           "size": 11.5, "font": MONO, "color": MUT, "space": 14, "lead": 1.5}],
         h=Inches(1.6))

    rx = ML + lw + Inches(0.5)
    rw = CW - lw - Inches(0.5)
    text(s, rx, y, rw,
         [{"t": "Those records are exactly the edges we model today.",
           "size": 19, "bold": True, "lead": 1.3},
          {"t": "Feed them into the same engine and the modeled links become "
                "real provenance. No new algorithm — the same aggregation on "
                "real edges.",
           "size": 13, "font": MONO, "color": MUT, "space": 12, "lead": 1.5},
          {"t": "Wedge: recall response.  Moat: the traceability system of "
                "record the law requires.",
           "size": 12.5, "font": MONO, "color": CYAN, "space": 14, "lead": 1.5}],
         h=Inches(2.1))

    yb = Inches(5.42)
    text(s, ML, yb, CW,
         [{"t": "Who buys it", "size": 10, "font": MONO, "color": MUT,
           "caps": True, "track": 1.8}], h=Inches(0.3))
    buyers = ["Regional restaurant chains", "Ghost-kitchen groups",
              "Distributors", "Grocers"]
    bwid = int((CW - Inches(0.9)) / 4)
    x = ML
    for b in buyers:
        rect(s, x, yb + Inches(0.38), Emu(bwid), Inches(0.62), fill=PANEL, edge=LINE)
        text(s, x + Inches(0.16), yb + Inches(0.54), Emu(bwid) - Inches(0.32),
             [{"t": b, "size": 12, "color": INK}], h=Inches(0.4))
        x += bwid + Inches(0.3)

    text(s, ML, Inches(6.72), CW,
         [{"t": "~$30B by 2030   ·   food-traceability technology market "
                "(Grand View, 8.8% CAGR; other houses put 2030 at $10B–$44B)",
           "size": 12, "font": MONO, "color": MUT}], h=Inches(0.5))
    return s


def s14_ask(prs):
    s = new_slide(prs)
    chrome(s, "Business", 14)
    text(s, ML, Inches(2.20), CW,
         [{"t": "The next real milestone", "size": 11, "font": MONO,
           "color": MUT, "caps": True, "track": 2.0, "align": PP_ALIGN.CENTER},
          {"t": "One restaurant group.\nTheir real purchase records.\n"
                "Twelve months.",
           "size": 46, "bold": True, "space": 22, "lead": 1.25,
           "align": PP_ALIGN.CENTER}], h=Inches(3.2))
    rule(s, int((W - Inches(1.4)) / 2), Inches(5.78), Inches(1.4), LINE)
    text(s, ML, Inches(6.10), CW,
         [{"t": "That one step converts the modeled edges into real provenance "
                "and proves the platform hinge on live data.\nNo hockey stick "
                "— one operator, their records, on their own box. Same engine "
                "the whole way up.",
           "size": 13, "font": MONO, "color": MUT, "lead": 1.6,
           "align": PP_ALIGN.CENTER}], h=Inches(1.2))
    return s


def s15_internal(prs):
    s = new_slide(prs)
    y = chrome(s, "Internal — do not present", 15,
               "The six we never overclaim")

    items = [
        ("Recall responder, not outbreak detector",
         "~100% of our feed (FDA lot-coded recalls); only a minority of real epidemiological outbreaks."),
        ("FSIS is a labeled blind spot",
         "FDA-only. Boar's Head 2024 wouldn't appear. Pre-empt it; don't get caught."),
        ("Local inference is currently thin",
         "The LLM writes one narration sentence; deterministic Mongo makes every decision."),
        ("Change stream isn't wired to run() yet",
         "The Watcher observes inserts; it does not auto-trigger the pipeline. We trigger runs explicitly."),
        ("Supplier→site links are modeled",
         "The FSMA-204 gap we exist to close — labeled modeled, not passed off as ground truth."),
        ("Telegram is draft-not-send",
         "The Comms output is prepared, not auto-dispatched to a live operator channel."),
    ]
    cy = y
    for i, (head, body) in enumerate(items, 1):
        text(s, ML, cy, Inches(0.5),
             [{"t": f"{i}", "size": 15, "bold": True, "color": CRIT,
               "font": MONO}], h=Inches(0.4))
        text(s, ML + Inches(0.5), cy, CW - Inches(0.5),
             [{"t": head, "size": 14, "bold": True},
              {"t": body, "size": 11, "font": MONO, "color": MUT,
               "space": 4, "lead": 1.35}], h=Inches(0.72))
        cy += Inches(0.76)
    return s


# --------------------------------------------------------------------------
def load_notes():
    """Pull the 15 note blocks straight out of the Figma script, so the two
    decks cannot drift apart."""
    import json
    src = (ROOT / "scripts" / "figma-speaker-notes.js").read_text()
    body = src[src.index("const N=[") + len("const N=["):]
    body = body[:body.index("\n];")]
    notes = []
    for raw in body.split("\n"):
        raw = raw.strip()
        if raw.startswith('"') and raw.endswith('",'):
            notes.append(json.loads(raw[:-1]))
    return notes


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    builders = [s01_title, s02_today, s03_contrast, s04_local, s05_demo,
                s06_recall, s07_trace, s08_signals, s09_output, s10_coverage,
                s11_instrument, s12_swarm, s13_business, s14_ask, s15_internal]

    notes = load_notes()
    if len(notes) != len(builders):
        raise SystemExit(f"note count {len(notes)} != slide count {len(builders)}")

    for build, note in zip(builders, notes):
        slide = build(prs)
        slide.notes_slide.notes_text_frame.text = note

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"{OUT.relative_to(ROOT)}  {OUT.stat().st_size:,} bytes")
    print(f"{len(builders)} slides, {len(notes)} speaker notes, "
          f"site table = {SITE_TABLE}")


if __name__ == "__main__":
    main()
